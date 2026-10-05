# -*- coding: utf-8 -*-
"""
Prova detalhada — Excel ↓ cols + Revisar dados (`CREDITO-SCORE-XLSX-COLS`).

  set AGRO_PIN_TESTE=9973
  set PYTHONIOENCODING=utf-8
  python scripts/verify_credito_score_xlsx_path.py

HTTP via Django test Client (não precisa runserver).
Snapshots de prova: savepoint + rollback (não grava na loja).
"""
from __future__ import annotations

import ast
import io
import os
import re
import sys
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")

import django

django.setup()

from django.contrib.auth import get_user_model
from django.db import transaction
from django.test import Client, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from produtos.caixa_util import validar_pin_operador
from produtos.credito_score_views import (
    _listar_snapshots_filtrados,
    _montar_xlsx_laboratorio,
)
from produtos.credito_score_shadow import rotulo_candidato_revisao
from produtos.models import ClienteAgro, ClienteAnaliseCreditoAgro

PIN = (os.environ.get("AGRO_PIN_TESTE") or "9973").strip()
oks: list[str] = []
fails: list[str] = []

HEADERS_ESPERADOS = [
    "ID cliente",
    "Cliente",
    "Score",
    "Classificação",
    "Confiança",
    "Limite cadastrado",
    "Limite atual (efetivo)",
    "Em aberto",
    "Saldo vencido",
    "Média fiado 3m",
    "Limite sugerido",
    "Dif. sugerido − atual",
    "Situação",
    "Candidato revisão",
    "Qtd títulos analisados",
    "Qtd quitada",
    "Qtd vencida",
    "% pago em dia",
    "Pts pontualidade",
    "Pts situação atual",
    "Pts quitação",
    "Pts frequência",
    "Pts relacionamento",
    "Maior atraso (dias)",
    "Calculado em",
    "Alertas",
]


def check(name: str, cond: bool, detail: str = "") -> None:
    if cond:
        oks.append(name)
        print(f"  OK  {name}" + (f" — {detail}" if detail else ""))
    else:
        fails.append(name)
        print(f"  FAIL {name}" + (f" — {detail}" if detail else ""))


def _hosts() -> list[str]:
    from django.conf import settings

    hosts = list(getattr(settings, "ALLOWED_HOSTS", []) or [])
    for h in ("testserver", "localhost", "127.0.0.1"):
        if h not in hosts:
            hosts.append(h)
    return hosts


def test_contratos_arquivo() -> None:
    print("== 1) Contratos arquivo / rota / botão ==")
    urls = (ROOT / "produtos/urls.py").read_text(encoding="utf-8")
    views = (ROOT / "produtos/credito_score_views.py").read_text(encoding="utf-8")
    tpl = (ROOT / "produtos/templates/produtos/credito_score_laboratorio.html").read_text(
        encoding="utf-8"
    )
    dual = (ROOT / "produtos/static/produtos/js/agro_dual_window.js").read_text(
        encoding="utf-8"
    )

    check("rota_export", "fiado/analise-credito/export-xlsx/" in urls)
    check(
        "rota_name",
        "credito_score_laboratorio_export_xlsx" in urls,
    )
    check(
        "view_fn",
        "def credito_score_laboratorio_export_xlsx" in views,
    )
    check(
        "view_decorator",
        "@credito_score_shadow_required" in views
        and "credito_score_laboratorio_export_xlsx" in views,
    )
    check("view_require_get", "@require_GET" in views)
    check(
        "view_content_type",
        "spreadsheetml.sheet" in views and "Content-Disposition" in views,
    )
    check("view_usa_filtros", "_filtros_request" in views and "_listar_snapshots_filtrados" in views)
    check("view_openpyxl", "_montar_xlsx_laboratorio" in views and "Workbook" in views)

    # Isolamento: export não escreve limite/cliente/fiado
    tree = ast.parse(views)
    export_fn = None
    for node in tree.body:
        if isinstance(node, ast.FunctionDef) and node.name == "credito_score_laboratorio_export_xlsx":
            export_fn = node
            break
    check("ast_export_fn", export_fn is not None)
    if export_fn is not None:
        src = ast.get_source_segment(views, export_fn) or ""
        check("export_sem_save", ".save(" not in src and "objects.create" not in src)
        check("export_sem_update", ".update(" not in src and "objects.filter" not in src)
        check("export_sem_analisar", "analisar_cliente" not in src and "persistir_analise" not in src)

    check("tpl_botao_excel", "Excel ↓" in tpl or "Excel &darr;" in tpl or ">Excel" in tpl)
    check(
        "tpl_url_export",
        "credito_score_laboratorio_export_xlsx" in tpl,
    )
    check(
        "tpl_passa_filtros",
        "score_min={{ score_min|urlencode }}" in tpl and "classificacao={{ classificacao|urlencode }}" in tpl,
    )
    check("tpl_badge_revisar", "revisar_dados" in tpl and "Revisar dados" in tpl)
    check("tpl_badge_candidato", "candidato_revisao" in tpl and "Candidato revisão" in tpl)

    det = (
        ROOT / "produtos/templates/produtos/credito_score_cliente_detalhe.html"
    ).read_text(encoding="utf-8")
    check("det_badge_revisar", "revisar_dados" in det and "Revisar dados" in det)

    shadow = (ROOT / "produtos/credito_score_shadow.py").read_text(encoding="utf-8")
    check("motor_fn_alerta", "def alerta_inconsistencia_pagamento" in shadow)
    check("motor_fn_rotulo", "def rotulo_candidato_revisao" in shadow)
    check("motor_marca_baixas", "sem baixas suficientes" in shadow)
    check(
        "motor_gate_candidato",
        "revisar_dados = alerta_inconsistencia_pagamento(alertas)" in shadow
        and "candidato = False" in shadow,
    )
    check("motor_flag_json", "revisar_dados_inconsistencia" in shadow)
    check("score_peso_pontualidade_45", '* Decimal("45")' in shadow)
    check("score_peso_situacao_25", "situacao_pts = 25" in shadow)
    check("score_peso_quit_freq_rel_10", 'Decimal("10")' in shadow)
    check(
        "view_cols_indicadores",
        "titulos_analisados" in views
        and "pct_pago_em_dia" in views
        and "pts_pontualidade" in views
        and "pts_relacionamento" in views,
    )
    check("view_usa_rotulo", "rotulo_candidato_revisao" in views)

    check(
        "dual_export_coberto",
        "p.indexOf('/fiado/analise-credito/') === 0" in dual
        or 'p.indexOf("/fiado/analise-credito/") === 0' in dual,
    )

    # Fiado operacional / PDV não linkam export
    fiado_hits = 0
    for p in (ROOT / "produtos/templates/produtos").glob("*fiado*"):
        if "credito_score" in p.name:
            continue
        t = p.read_text(encoding="utf-8", errors="ignore")
        if "export-xlsx" in t or "credito_score_laboratorio_export" in t:
            fiado_hits += 1
    check("sem_link_fiado_operacional", fiado_hits == 0, str(fiado_hits))

    wiz = ROOT / "produtos/static/produtos/js/pdv_wizard.js"
    if wiz.is_file():
        w = wiz.read_text(encoding="utf-8")
        check("pdv_sem_export", "export-xlsx" not in w and "credito_score" not in w)


def test_pin() -> None:
    print("== 2) PIN operacional ==")
    ok, msg = validar_pin_operador(PIN)
    check("pin_9973", ok, (msg or "")[:60])


def test_builder_unitario() -> None:
    print("== 3) Builder openpyxl (unitário) ==")
    check(
        "rotulo_revisar_dados",
        rotulo_candidato_revisao(
            candidato=True,
            alertas=["Título quitado #9 sem baixas suficientes (pago R$ 10 · baixas R$ 0)."],
        )
        == "Revisar dados",
    )
    check(
        "rotulo_sim",
        rotulo_candidato_revisao(candidato=True, alertas=[]) == "Sim",
    )
    check(
        "rotulo_nao",
        rotulo_candidato_revisao(candidato=False, alertas=[]) == "Não",
    )
    rows = [
        {
            "cliente_pk": 77,
            "cliente_nome": "Prova XLSX Alfa",
            "score": 82,
            "classificacao_label": "Bom",
            "confianca_label": "Alta",
            "limite_cadastrado": Decimal("200.00"),
            "limite_efetivo": Decimal("200.00"),
            "saldo_aberto": Decimal("40.00"),
            "saldo_vencido": Decimal("0"),
            "media_fiado_3m": Decimal("100.50"),
            "limite_sugerido": Decimal("250.00"),
            "situacao_label": "Em dia",
            "candidato_revisao": True,
            "candidato_revisao_label": "Sim",
            "titulos_analisados": 5,
            "titulos_quitados": 4,
            "titulos_vencidos_qtd": 0,
            "pct_pago_em_dia": 75.0,
            "pts_pontualidade": 40,
            "pts_situacao": 25,
            "pts_quitacao": 10,
            "pts_frequencia": 8,
            "pts_relacionamento": 6,
            "maior_atraso_dias": 0,
            "calculado_em": timezone.now(),
            "alertas": ["prova"],
        }
    ]
    raw = _montar_xlsx_laboratorio(rows)
    check("xlsx_magic_pk", raw[:2] == b"PK", str(len(raw)))
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.active
    hdrs = [c.value for c in ws[1]]
    check("xlsx_headers", hdrs == HEADERS_ESPERADOS, str(hdrs[:5]))
    check("xlsx_id", ws.cell(2, 1).value == 77)
    check("xlsx_nome", ws.cell(2, 2).value == "Prova XLSX Alfa")
    check("xlsx_score", ws.cell(2, 3).value == 82)
    check("xlsx_media", abs(float(ws.cell(2, 10).value) - 100.50) < 0.001)
    check("xlsx_diff", abs(float(ws.cell(2, 12).value) - 50.0) < 0.001)
    check("xlsx_candidato", ws.cell(2, 14).value == "Sim")
    check("xlsx_qtd_analisados", ws.cell(2, 15).value == 5)
    check("xlsx_pct_em_dia", abs(float(ws.cell(2, 18).value) - 75.0) < 0.001)
    check("xlsx_pts_pont", ws.cell(2, 19).value == 40)
    check("xlsx_auto_filter", bool(ws.auto_filter.ref))
    check("xlsx_freeze", ws.freeze_panes == "A2")

    rows_rev = [
        {
            **rows[0],
            "candidato_revisao": False,
            "candidato_revisao_label": "Revisar dados",
            "alertas": ["Título quitado #1 sem baixas suficientes (x)."],
        }
    ]
    wb2 = load_workbook(io.BytesIO(_montar_xlsx_laboratorio(rows_rev)), data_only=True)
    check("xlsx_revisar_dados", wb2.active.cell(2, 14).value == "Revisar dados")


def test_http_e_filtros() -> None:
    print("== 4) HTTP gate + export + filtros ==")
    User = get_user_model()
    superu = User.objects.filter(is_superuser=True).order_by("pk").first()
    check("tem_superuser", superu is not None, getattr(superu, "username", ""))
    if not superu:
        return

    url_lab = reverse("credito_score_laboratorio")
    url_x = reverse("credito_score_laboratorio_export_xlsx")
    check("reverse_export", url_x.endswith("/fiado/analise-credito/export-xlsx/") or "export-xlsx" in url_x)
    c = Client()
    hosts = _hosts()

    with override_settings(
        AGRO_CREDITO_SCORE_SHADOW_ENABLED=False,
        AGRO_CREDITO_SCORE_SHADOW_USERNAMES="",
        ALLOWED_HOSTS=hosts,
    ):
        c.force_login(superu)
        check("flag_off_export_404", c.get(url_x).status_code == 404)

    with override_settings(
        AGRO_CREDITO_SCORE_SHADOW_ENABLED=True,
        AGRO_CREDITO_SCORE_SHADOW_USERNAMES="renan",
        ALLOWED_HOSTS=hosts,
    ):
        comum = (
            User.objects.filter(is_superuser=False, is_staff=False)
            .exclude(username__iexact="renan")
            .first()
        )
        if comum:
            c.force_login(comum)
            check("operador_export_404", c.get(url_x).status_code == 404, comum.username)

        c.force_login(superu)
        r_lab = c.get(url_lab)
        check("lab_200", r_lab.status_code == 200, str(r_lab.status_code))
        check("lab_tem_botao_excel", b"Excel" in r_lab.content and b"export-xlsx" in r_lab.content)

        # Snapshot de prova isolado
        with transaction.atomic():
            sid = transaction.savepoint()
            try:
                cli = ClienteAgro.objects.create(
                    nome="ZZ PROVA CREDITO XLSX NAO USAR",
                    limite_fiado_local=Decimal("300.00"),
                    ativo=True,
                )
                ClienteAnaliseCreditoAgro.objects.create(
                    cliente=cli,
                    regra_versao="shadow_v1",
                    score=91,
                    classificacao=ClienteAnaliseCreditoAgro.Classificacao.EXCELENTE,
                    confianca=ClienteAnaliseCreditoAgro.Confianca.ALTA,
                    limite_cadastrado_snapshot=Decimal("300.00"),
                    limite_efetivo_snapshot=Decimal("300.00"),
                    saldo_aberto_snapshot=Decimal("10.00"),
                    saldo_vencido_snapshot=Decimal("0"),
                    tem_vencido_snapshot=False,
                    media_fiado_3m=Decimal("120.00"),
                    limite_sugerido=Decimal("400.00"),
                    maior_atraso_dias=0,
                    indicadores_json={
                        "candidato_revisao_limite": True,
                        "titulos_analisados_janela": 8,
                        "titulos_quitados_avaliaveis": 6,
                        "titulos_vencidos_atualmente": 0,
                        "pagamentos_em_dia": 5,
                        "pontualidade_pontos": 40,
                        "situacao_atual_pontos": 25,
                        "quitacao_pontos": 9,
                        "frequencia_pontos": 8,
                        "relacionamento_pontos": 7,
                    },
                    alertas_json=["prova-xlsx"],
                )
                lim_antes = ClienteAgro.objects.get(pk=cli.pk).limite_fiado_local

                rows_lista = _listar_snapshots_filtrados({"q": "ZZ PROVA CREDITO XLSX"})
                check("filtro_q_lista", len(rows_lista) == 1, str(len(rows_lista)))

                r = c.get(url_x, {"q": "ZZ PROVA CREDITO XLSX"})
                check("export_200", r.status_code == 200, str(r.status_code))
                ctype = r.get("Content-Type", "")
                check("export_ctype", "spreadsheetml.sheet" in ctype, ctype[:60])
                disp = r.get("Content-Disposition", "")
                check("export_disposition", "attachment" in disp and ".xlsx" in disp, disp[:80])
                body = r.content
                check("export_magic", body[:2] == b"PK", str(len(body)))

                wb = load_workbook(io.BytesIO(body), data_only=True)
                ws = wb.active
                hdrs = [cell.value for cell in ws[1]]
                check("export_headers", hdrs == HEADERS_ESPERADOS)
                # 1 header + 1 data (filtro q)
                check("export_1_linha", ws.max_row == 2, f"max_row={ws.max_row}")
                check("export_id_match", int(ws.cell(2, 1).value) == cli.pk)
                check("export_nome_match", "ZZ PROVA CREDITO XLSX" in str(ws.cell(2, 2).value))
                check("export_score_match", int(ws.cell(2, 3).value) == 91)
                check(
                    "export_sugerido",
                    abs(float(ws.cell(2, 11).value) - 400.0) < 0.001,
                )
                check(
                    "export_diff",
                    abs(float(ws.cell(2, 12).value) - 100.0) < 0.001,
                )
                check("export_candidato", ws.cell(2, 14).value == "Sim")
                check("export_qtd_analisados", ws.cell(2, 15).value == 8)
                check("export_qtd_quitada", ws.cell(2, 16).value == 6)
                check(
                    "export_pct_em_dia",
                    abs(float(ws.cell(2, 18).value) - round(100.0 * 5 / 6, 2)) < 0.01,
                )
                check("export_pts_pont", ws.cell(2, 19).value == 40)
                check("export_pts_sit", ws.cell(2, 20).value == 25)
                check("export_pts_quit", ws.cell(2, 21).value == 9)
                check("export_pts_freq", ws.cell(2, 22).value == 8)
                check("export_pts_rel", ws.cell(2, 23).value == 7)

                # Regra provisória: flag candidato True + alerta baixas → Revisar dados
                cli2 = ClienteAgro.objects.create(
                    nome="ZZ PROVA CREDITO REVISAR DADOS NAO USAR",
                    limite_fiado_local=Decimal("500.00"),
                    ativo=True,
                )
                ClienteAnaliseCreditoAgro.objects.create(
                    cliente=cli2,
                    regra_versao="shadow_v1",
                    score=90,
                    classificacao=ClienteAnaliseCreditoAgro.Classificacao.EXCELENTE,
                    confianca=ClienteAnaliseCreditoAgro.Confianca.ALTA,
                    limite_cadastrado_snapshot=Decimal("500.00"),
                    limite_efetivo_snapshot=Decimal("500.00"),
                    saldo_aberto_snapshot=Decimal("0"),
                    saldo_vencido_snapshot=Decimal("0"),
                    tem_vencido_snapshot=False,
                    media_fiado_3m=Decimal("200.00"),
                    limite_sugerido=Decimal("600.00"),
                    maior_atraso_dias=0,
                    indicadores_json={"candidato_revisao_limite": True},
                    alertas_json=[
                        "Título quitado #99 sem baixas suficientes "
                        "(pago no título R$ 100 · baixas R$ 0)."
                    ],
                )
                rows_rev = _listar_snapshots_filtrados({"q": "ZZ PROVA CREDITO REVISAR"})
                check(
                    "lista_revisar_label",
                    len(rows_rev) == 1 and rows_rev[0].get("revisar_dados") is True,
                )
                check(
                    "lista_nao_candidato",
                    rows_rev[0].get("candidato_revisao") is False
                    and rows_rev[0].get("candidato_revisao_label") == "Revisar dados",
                )
                r_rev = c.get(url_x, {"q": "ZZ PROVA CREDITO REVISAR"})
                wb_rev = load_workbook(io.BytesIO(r_rev.content), data_only=True)
                check("export_revisar_dados", wb_rev.active.cell(2, 14).value == "Revisar dados")

                r_lab2 = c.get(url_lab, {"q": "ZZ PROVA CREDITO REVISAR"})
                check("lab_html_revisar", b"Revisar dados" in r_lab2.content)

                # Filtro score_min alto → 0 linhas de dados
                r0 = c.get(url_x, {"q": "ZZ PROVA CREDITO XLSX", "score_min": "99"})
                check("export_filtro_score_200", r0.status_code == 200)
                wb0 = load_workbook(io.BytesIO(r0.content), data_only=True)
                check(
                    "export_filtro_score_vazio",
                    wb0.active.max_row == 1,
                    f"max_row={wb0.active.max_row}",
                )

                # Contagem lista == planilha (sem filtro)
                rows_all = _listar_snapshots_filtrados({})
                r_all = c.get(url_x)
                wb_all = load_workbook(io.BytesIO(r_all.content), data_only=True)
                check(
                    "export_conta_igual_lista",
                    (wb_all.active.max_row - 1) == len(rows_all),
                    f"xlsx={wb_all.active.max_row - 1} lista={len(rows_all)}",
                )

                lim_depois = ClienteAgro.objects.get(pk=cli.pk).limite_fiado_local
                check("limite_cliente_inalterado", lim_antes == lim_depois, str(lim_depois))
            finally:
                transaction.savepoint_rollback(sid)

        # Limpa sobra de provas anteriores (se houver) + confirma rollback
        ClienteAgro.objects.filter(nome__icontains="ZZ PROVA CREDITO XLSX").delete()
        ClienteAgro.objects.filter(nome__icontains="ZZ PROVA CREDITO REVISAR").delete()
        check(
            "rollback_sem_lixo",
            not ClienteAgro.objects.filter(nome__icontains="ZZ PROVA CREDITO XLSX").exists()
            and not ClienteAgro.objects.filter(nome__icontains="ZZ PROVA CREDITO REVISAR").exists(),
        )


def test_regressao_leve() -> None:
    print("== 5) Regressão leve (shadow path) ==")
    import subprocess

    env = os.environ.copy()
    env["AGRO_PIN_TESTE"] = PIN
    env["PYTHONIOENCODING"] = "utf-8"
    r = subprocess.run(
        [sys.executable, str(ROOT / "scripts/verify_credito_score_shadow_path.py")],
        cwd=str(ROOT),
        env=env,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    out = (r.stdout or "") + (r.stderr or "")
    m = re.search(r"(\d+)\s+OK\s*[·.]\s*(\d+)\s+FAIL", out)
    if m:
        check(
            "shadow_path_ok",
            r.returncode == 0 and int(m.group(2)) == 0,
            f"{m.group(1)} OK · {m.group(2)} FAIL",
        )
    else:
        check("shadow_path_ok", r.returncode == 0, out[-220:].replace("\n", " "))


def main() -> int:
    print(f"=== verify CREDITO-SCORE-XLSX-COLS · PIN={PIN} ===")
    test_contratos_arquivo()
    test_pin()
    test_builder_unitario()
    test_http_e_filtros()
    test_regressao_leve()
    print(f"=== {len(oks)} OK · {len(fails)} FAIL ===")
    if fails:
        print("FAILS:", ", ".join(fails))
        print("PREP_FAILS=" + str(len(fails)))
        return 1
    print("PREP_FAILS=0")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
