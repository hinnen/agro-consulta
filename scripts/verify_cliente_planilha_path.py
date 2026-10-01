# -*- coding: utf-8 -*-
"""Prova — Excel clientes (`CLIENTE-XLSX-FIADO`). python scripts/verify_cliente_planilha_path.py"""
from __future__ import annotations

import os
import sys
import tempfile
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")

import django

django.setup()

from django.contrib.auth import get_user_model
from django.test import Client

from produtos.cliente_planilha_util import (
    COL_FIADO_USADO,
    COL_ID,
    COL_LIMITE_FIADO,
    COL_MEDIA_FIADO_3M,
    COL_MES_MAIS_FIADO,
    COL_NOME,
    COL_QTD_FIADO_3M,
    COL_VALOR_MES_ANTERIOR,
    COL_VALOR_MES_MAIS,
    COL_WHATSAPP,
    EXPORT_ONLY,
    IMPORT_EDIT_KEYS,
    _inicio_janela_fiado_3_meses,
    _label_mes,
    _meses_calendario_ultimos_3,
    aplicar_importacao_clientes,
    coletar_linhas_export_clientes,
    montar_xlsx_clientes,
    preview_importacao_clientes,
)
from datetime import date
from produtos.models import ClienteAgro, VendaAgro

fails: list[str] = []
oks: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    if cond:
        oks.append(name)
        print(f"  OK  {name}" + (f" — {detail}" if detail else ""))
    else:
        fails.append(name)
        print(f"  FAIL {name}" + (f" — {detail}" if detail else ""))


def test_media_mensal_calendario() -> None:
    print("== Média fiado/mês (3 meses calendário) ==")
    ref = date(2026, 10, 15)
    check("meses_ref", _meses_calendario_ultimos_3(ref) == [(2026, 10), (2026, 9), (2026, 8)])
    check("inicio_janela", _inicio_janela_fiado_3_meses(ref) == date(2026, 8, 1))
    util = (ROOT / "produtos/cliente_planilha_util.py").read_text(encoding="utf-8")
    check("media_div_3", "total_3m / Decimal(\"3\")" in util)
    check("help_media_mes", "média fiado/mês" in util.lower() or "media fiado/mes" in util.lower())


def test_arquivos() -> None:
    print("== Contratos UI / URLs ==")
    html = (ROOT / "produtos/templates/produtos/clientes_lista.html").read_text(encoding="utf-8")
    urls = (ROOT / "produtos/urls.py").read_text(encoding="utf-8")
    util = (ROOT / "produtos/cliente_planilha_util.py").read_text(encoding="utf-8")
    views = (ROOT / "produtos/views_cliente_planilha.py").read_text(encoding="utf-8")

    check("btn_export", 'id="clientes-btn-export-xlsx"' in html and "Excel ↓" in html)
    check("btn_import", 'id="clientes-btn-import-xlsx"' in html and "Excel ↑" in html)
    check("modal_import", 'id="clientes-import-modal"' in html)
    check("previa_btn", 'id="clientes-import-previa-btn"' in html)
    check("confirmar_btn", 'id="clientes-import-confirmar"' in html)
    check("url_export", "api/clientes/export-xlsx/" in urls)
    check("url_preview", "api/clientes/import-preview/" in urls)
    check("url_aplicar", "api/clientes/import-aplicar/" in urls)
    check("view_export", "def api_clientes_export_xlsx" in views)
    check("col_media_3m", COL_MEDIA_FIADO_3M in util)
    check("col_mes_mais", COL_MES_MAIS_FIADO in util)
    check("export_only_fiado", COL_MEDIA_FIADO_3M in EXPORT_ONLY and COL_MES_MAIS_FIADO in EXPORT_ONLY)
    check("valores_so_leitura", COL_VALOR_MES_MAIS in EXPORT_ONLY and COL_VALOR_MES_ANTERIOR in EXPORT_ONLY)
    check("valores_nao_editaveis", COL_VALOR_MES_MAIS not in IMPORT_EDIT_KEYS and COL_VALOR_MES_ANTERIOR not in IMPORT_EDIT_KEYS)
    check("limite_editavel", COL_LIMITE_FIADO in IMPORT_EDIT_KEYS)
    check("aba_como_usar", '"Como usar"' in util or "'Como usar'" in util)
    check("analise_importacao", "def analise_importacao_clientes" in util)
    check("aplicar_usa_analise", "analise_importacao_clientes(path)" in util and "alteracoes[:400]" not in util.split("def aplicar_importacao_clientes")[1].split("def ")[0])
    check("help_limite_zero_pdv", "PDV mostra o padrão" in util or "padrão (R$ 5.000)" in util)
    check("help_whatsapp_vazio_apaga", "exceto WhatsApp" in util and "apaga o número" in util)
    check("patch_whatsapp_vazio", "patch[COL_WHATSAPP] = v_wa[:20] if v_wa else \"\"" in util)


def test_export_xlsx_bytes() -> None:
    print("== Export XLSX ==")
    rows = coletar_linhas_export_clientes()[:5]
    if not rows:
        cli = ClienteAgro.objects.create(nome="Cliente Prova Planilha")
        rows = coletar_linhas_export_clientes()[:5]
        cli.delete()
    data = montar_xlsx_clientes(rows)
    check("xlsx_bytes", len(data) > 2000, f"{len(data)} bytes")
    tmp = Path(tempfile.mkdtemp()) / "t.xlsx"
    tmp.write_bytes(data)
    from openpyxl import load_workbook

    wb = load_workbook(tmp)
    ws = wb["Clientes"]
    check("sheet_clientes", ws.title == "Clientes")
    check("header_nome", ws.cell(1, 3).value in ("Nome",) or "Nome" in str(ws.cell(1, 1).value))
    hdrs = [ws.cell(1, c).value for c in range(1, 30)]
    check("hdr_media", "Média fiado/mês (3 meses)" in hdrs)
    check("hdr_mes", "Mês que mais comprou fiado" in hdrs)
    check("hdr_valor_mes", "Valor do mês que mais comprou" in hdrs)
    check("hdr_valor_ant", "Valor mês anterior" in hdrs)
    check("hdr_limite", "Limite fiado" in hdrs)
    check("aba_instr", "Como usar" in wb.sheetnames)
    wb.close()
    tmp.unlink(missing_ok=True)


def test_valores_mes() -> None:
    """Mês com mais compras = contagem. Valor = soma desse mês. Mês anterior = calendário local."""
    print("== Valores de mês ==")
    from datetime import datetime, timedelta
    from zoneinfo import ZoneInfo

    from django.utils import timezone

    tz = ZoneInfo("America/Sao_Paulo")
    hoje = timezone.localdate()
    ant = hoje.replace(day=1) - timedelta(days=1)
    cli = ClienteAgro.objects.create(nome="ZZ Prova Mes Valor", limite_fiado_local=Decimal("0"))
    pks: list[int] = []
    try:
        def _venda(quando: datetime, valor: str, *, fiado: bool = True, devolvida=None) -> None:
            v = VendaAgro.objects.create(
                cliente_nome=cli.nome,
                cliente_id_erp=f"agro:{cli.pk}",
                total=Decimal(valor),
                forma_pagamento="Fiado" if fiado else "Dinheiro",
                pagamentos_json=[{"formaPagamento": "Fiado" if fiado else "Dinheiro", "valorPagamento": valor}],
            )
            VendaAgro.objects.filter(pk=v.pk).update(criado_em=quando, devolvida_em=devolvida)
            pks.append(v.pk)

        # 23h30 do último dia do mês passado = dia seguinte em UTC. Tem que cair no mês passado.
        _venda(datetime(ant.year, ant.month, ant.day, 23, 30, tzinfo=tz), "300.00")
        _venda(datetime(hoje.year, hoje.month, hoje.day, 10, 0, tzinfo=tz), "10.00")
        _venda(datetime(hoje.year, hoje.month, hoje.day, 11, 0, tzinfo=tz), "20.00")
        _venda(datetime(hoje.year, hoje.month, hoje.day, 12, 0, tzinfo=tz), "999.00", fiado=False)
        _venda(
            datetime(hoje.year, hoje.month, hoje.day, 13, 0, tzinfo=tz),
            "500.00",
            devolvida=timezone.now(),
        )

        rows = coletar_linhas_export_clientes()
        row = next(r for r in rows if r[COL_ID] == cli.pk)
        check("qtd_so_fiado", row[COL_QTD_FIADO_3M] == 3, str(row[COL_QTD_FIADO_3M]))
        check("mes_e_o_atual", row[COL_MES_MAIS_FIADO] == _label_mes(hoje.year, hoje.month), row[COL_MES_MAIS_FIADO])
        check("valor_do_mes_campeao", abs(float(row[COL_VALOR_MES_MAIS]) - 30.0) < 0.001, str(row[COL_VALOR_MES_MAIS]))
        check("valor_mes_anterior", abs(float(row[COL_VALOR_MES_ANTERIOR]) - 300.0) < 0.001, str(row[COL_VALOR_MES_ANTERIOR]))
        # Mês atual 30 + mês anterior 300 + terceiro mês 0 → média mensal 110 (não média por compra).
        check("media_por_mes", abs(float(row[COL_MEDIA_FIADO_3M]) - 110.0) < 0.001, str(row[COL_MEDIA_FIADO_3M]))
        check("aberto_nao_mexe", abs(float(row[COL_FIADO_USADO]) - 0.0) < 0.001, str(row[COL_FIADO_USADO]))

        row[COL_VALOR_MES_MAIS] = 99999
        row[COL_VALOR_MES_ANTERIOR] = 1
        data = montar_xlsx_clientes([row])
        tmp = Path(tempfile.mkdtemp()) / "mes.xlsx"
        tmp.write_bytes(data)
        from openpyxl import load_workbook

        wb = load_workbook(tmp)
        ws = wb["Clientes"]
        hdrs = [ws.cell(1, c).value for c in range(1, 30)]
        i_mes = hdrs.index("Valor do mês que mais comprou") + 1
        i_ant = hdrs.index("Valor mês anterior") + 1
        check("cinza_valor_mes", (ws.cell(1, i_mes).fill.fgColor.rgb or "").endswith("F1F5F9"))
        check("cinza_valor_ant", (ws.cell(1, i_ant).fill.fgColor.rgb or "").endswith("F1F5F9"))
        wb.close()
        prev = preview_importacao_clientes(tmp)
        campos = [c.get("campo") for a in prev.get("alteracoes", []) for c in a.get("campos", [])]
        check("import_ignora_valores", COL_VALOR_MES_MAIS not in campos and COL_VALOR_MES_ANTERIOR not in campos, str(campos))
        tmp.unlink(missing_ok=True)
    finally:
        VendaAgro.objects.filter(pk__in=pks).delete()
        cli.delete()


def test_import_roundtrip() -> None:
    print("== Import prévia + aplicar ==")
    cli = ClienteAgro.objects.create(
        nome="ZZ Prova Excel Cliente",
        whatsapp="11999990001",
        limite_fiado_local=Decimal("0"),
    )
    try:
        rows = coletar_linhas_export_clientes()
        row = next(r for r in rows if r[COL_ID] == cli.pk)
        row[COL_LIMITE_FIADO] = 150.0
        row[COL_NOME] = "ZZ Prova Excel Cliente Editado"
        data = montar_xlsx_clientes([row])
        tmp = Path(tempfile.mkdtemp()) / "imp.xlsx"
        tmp.write_bytes(data)

        prev = preview_importacao_clientes(tmp)
        check("prev_ok", prev.get("n_alteracoes", 0) >= 1, f"alt={prev.get('n_alteracoes')}")
        check("prev_limite", any(
            c.get("campo") == COL_LIMITE_FIADO for a in prev.get("alteracoes", []) for c in a.get("campos", [])
        ))

        User = get_user_model()
        user = User.objects.filter(is_superuser=True).first() or User.objects.first()
        r = aplicar_importacao_clientes(tmp, user, nome_arquivo="prova.xlsx")
        check("apply_ok", r.get("clientes_alterados", 0) >= 1)

        cli.refresh_from_db()
        check("nome_gravado", cli.nome == "ZZ Prova Excel Cliente Editado")
        check("limite_gravado", Decimal(str(cli.limite_fiado_local)) == Decimal("150.00"))
        check("editado_local", cli.editado_local is True)
        tmp.unlink(missing_ok=True)
    finally:
        cli.delete()


def test_import_whatsapp_vazio_limpa() -> None:
    print("== Import WhatsApp vazio apaga ==")
    cli = ClienteAgro.objects.create(
        nome="ZZ Prova WA Limpar",
        whatsapp="11988887777",
        limite_fiado_local=Decimal("0"),
    )
    try:
        rows = coletar_linhas_export_clientes()
        row = next(r for r in rows if r[COL_ID] == cli.pk)
        row[COL_WHATSAPP] = ""
        data = montar_xlsx_clientes([row])
        tmp = Path(tempfile.mkdtemp()) / "wa_clear.xlsx"
        tmp.write_bytes(data)

        prev = preview_importacao_clientes(tmp)
        wa_campos = [
            c
            for a in prev.get("alteracoes", [])
            for c in a.get("campos", [])
            if c.get("campo") == COL_WHATSAPP
        ]
        check("prev_whatsapp", len(wa_campos) == 1, str(wa_campos))
        check("prev_para_vazio", wa_campos and wa_campos[0].get("para") == "", str(wa_campos))

        User = get_user_model()
        user = User.objects.filter(is_superuser=True).first() or User.objects.first()
        aplicar_importacao_clientes(tmp, user, nome_arquivo="wa_clear.xlsx")
        cli.refresh_from_db()
        check("whatsapp_gravado_vazio", cli.whatsapp == "")
        tmp.unlink(missing_ok=True)
    finally:
        cli.delete()


def test_import_alem_de_400_e_nao_mexe_fiado() -> None:
    """Gravação aplica todas as linhas; fiado em aberto não entra no patch."""
    print("== Import >400 + fiado só leitura ==")
    from decimal import Decimal
    from unittest.mock import MagicMock, patch

    import django

    django.setup()
    from openpyxl import Workbook

    from produtos.cliente_planilha_util import (
        COL_FIADO_USADO,
        COL_LIMITE_FIADO,
        EXPORT_ONLY,
        IMPORT_EDIT_KEYS,
        aplicar_importacao_clientes,
        preview_importacao_clientes,
    )

    check("fiado_aberto_so_leitura", COL_FIADO_USADO in EXPORT_ONLY and COL_FIADO_USADO not in IMPORT_EDIT_KEYS)
    n = 401
    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "Limite fiado", "Fiado em aberto agora"])
    for i in range(1, n + 1):
        ws.append([i, 0.01, 99999])
    tmp = Path(tempfile.mkdtemp()) / "muitos.xlsx"
    wb.save(tmp)

    clis: dict[int, MagicMock] = {}

    def _filter(pk):
        qs = MagicMock()
        cli = clis.get(pk)
        if cli is None:
            cli = MagicMock()
            cli.pk = pk
            cli.nome = f"C{pk}"
            cli.whatsapp = ""
            cli.cpf = ""
            cli.ativo = True
            cli.cep = cli.uf = cli.cidade = cli.bairro = ""
            cli.logradouro = cli.numero = cli.complemento = ""
            cli.plus_code = cli.referencia_rural = cli.maps_url_manual = ""
            cli.saldo_cashback = Decimal("0")
            cli.saldo_vale_credito = Decimal("0")
            cli.limite_fiado_local = Decimal("0")
            clis[pk] = cli
        qs.first.return_value = cli
        return qs

    limites: list[tuple] = []

    def _def_limite(pk, valor, usuario=""):
        limites.append((pk, Decimal(str(valor))))
        return clis[pk]

    hist = MagicMock()
    hist.pk = 1
    hist.n_campos = n

    with (
        patch("produtos.cliente_planilha_util.ClienteAgro") as CM,
        patch("produtos.fiado_gestao_util.definir_limite_fiado_cliente", side_effect=_def_limite),
        patch("produtos.cliente_planilha_util.CadastroPlanilhaImportHistoricoAgro") as HM,
    ):
        CM.objects.filter.side_effect = lambda **kw: _filter(kw["pk"])
        HM.objects.create.return_value = hist
        HM.Tipo.CADASTRO = "cadastro"
        prev = preview_importacao_clientes(tmp)
        check("preview_conta_401", prev.get("n_alteracoes") == n, str(prev.get("n_alteracoes")))
        check("preview_amostra_400", len(prev.get("alteracoes") or []) == 400)
        check("preview_avisa_corte", prev.get("alteracoes_preview_cortada") is True)
        user = MagicMock()
        user.is_authenticated = True
        user.get_username.return_value = "prova"
        r = aplicar_importacao_clientes(tmp, user, nome_arquivo="muitos.xlsx")
    check("gravou_401", r.get("clientes_alterados") == n, str(r.get("clientes_alterados")))
    check("chamou_limite_401", len(limites) == n, str(len(limites)))
    check("ultimo_e_401", limites and limites[-1][0] == n and limites[-1][1] == Decimal("0.01"))
    check("primeiro_e_1", limites and limites[0][0] == 1)
    tmp.unlink(missing_ok=True)


def test_http_endpoints() -> None:
    print("== HTTP (login) ==")
    from django.conf import settings

    if "testserver" not in settings.ALLOWED_HOSTS:
        settings.ALLOWED_HOSTS = [*settings.ALLOWED_HOSTS, "testserver"]
    User = get_user_model()
    user = User.objects.filter(is_superuser=True).first()
    if not user:
        check("http_skip", True, "sem user")
        return
    c = Client()
    c.force_login(user)
    r = c.get("/api/clientes/export-xlsx/")
    check("http_export_200", r.status_code == 200, str(r.status_code))
    check(
        "http_export_xlsx",
        "spreadsheetml" in (r.get("Content-Type") or ""),
        r.get("Content-Type", ""),
    )
    r2 = c.post("/api/clientes/import-preview/")
    check("http_preview_400_sem_arquivo", r2.status_code == 400)


def main() -> int:
    test_media_mensal_calendario()
    test_arquivos()
    test_export_xlsx_bytes()
    test_valores_mes()
    test_import_roundtrip()
    test_import_whatsapp_vazio_limpa()
    test_import_alem_de_400_e_nao_mexe_fiado()
    test_http_endpoints()
    print(f"\n== {len(oks)} OK · {len(fails)} FAIL ==")
    if fails:
        for f in fails:
            print("  -", f)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
