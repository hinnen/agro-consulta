"""Excel de clientes — exportar / importar (tela /clientes/, igual cadastro produtos)."""

from __future__ import annotations

import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter

from produtos.fiado_credito_util import (
    cliente_agro_pk_de_ref,
    resolver_cliente_fiado,
    valor_fiado_venda_local,
)
from produtos.models import CadastroPlanilhaImportHistoricoAgro, ClienteAgro, FiadoTituloAgro, VendaAgro

EXPORT_MAX_ROWS = 20000
IMPORT_MAX_ROWS = 5000
TIPO_HISTORICO = "clientes"
MESES_PT = ("Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez")

COL_ID = "id"
COL_CODIGO_ERP = "codigo_erp"
COL_NOME = "nome"
COL_WHATSAPP = "whatsapp"
COL_CPF = "cpf"
COL_ATIVO = "ativo"
COL_CEP = "cep"
COL_UF = "uf"
COL_CIDADE = "cidade"
COL_BAIRRO = "bairro"
COL_LOGRADOURO = "logradouro"
COL_NUMERO = "numero"
COL_COMPLEMENTO = "complemento"
COL_PLUS_CODE = "plus_code"
COL_REF_RURAL = "referencia_rural"
COL_MAPS_URL = "maps_url_manual"
COL_ENDERECO = "endereco"
COL_SALDO_CASHBACK = "saldo_cashback"
COL_SALDO_VALE = "saldo_vale_credito"
COL_LIMITE_FIADO = "limite_fiado_local"
COL_QTD_FIADO_3M = "qtd_compras_fiado_3m"
COL_MEDIA_FIADO_3M = "media_fiado_3m"
COL_MES_MAIS_FIADO = "mes_mais_comprou_fiado"
COL_VALOR_MES_MAIS = "valor_mes_mais_comprou"
COL_VALOR_MES_ANTERIOR = "valor_mes_anterior"
COL_FIADO_USADO = "fiado_usado_agora"

EXPORT_HEADERS: list[tuple[str, str]] = [
    ("ID", COL_ID),
    ("Código ERP", COL_CODIGO_ERP),
    ("Nome", COL_NOME),
    ("WhatsApp", COL_WHATSAPP),
    ("CPF", COL_CPF),
    ("Ativo", COL_ATIVO),
    ("CEP", COL_CEP),
    ("UF", COL_UF),
    ("Cidade", COL_CIDADE),
    ("Bairro", COL_BAIRRO),
    ("Logradouro", COL_LOGRADOURO),
    ("Número", COL_NUMERO),
    ("Complemento", COL_COMPLEMENTO),
    ("Plus Code", COL_PLUS_CODE),
    ("Referência rural", COL_REF_RURAL),
    ("Link Maps", COL_MAPS_URL),
    ("Endereço (resumo)", COL_ENDERECO),
    ("Saldo cashback", COL_SALDO_CASHBACK),
    ("Saldo vale", COL_SALDO_VALE),
    ("Limite fiado", COL_LIMITE_FIADO),
    ("Qtd compras fiado (3 meses)", COL_QTD_FIADO_3M),
    ("Média fiado (3 meses)", COL_MEDIA_FIADO_3M),
    ("Mês que mais comprou fiado", COL_MES_MAIS_FIADO),
    ("Valor do mês que mais comprou", COL_VALOR_MES_MAIS),
    ("Valor mês anterior", COL_VALOR_MES_ANTERIOR),
    ("Fiado em aberto agora", COL_FIADO_USADO),
]

EXPORT_ONLY = frozenset(
    {
        COL_ID,
        COL_CODIGO_ERP,
        COL_ENDERECO,
        COL_QTD_FIADO_3M,
        COL_MEDIA_FIADO_3M,
        COL_MES_MAIS_FIADO,
        COL_VALOR_MES_MAIS,
        COL_VALOR_MES_ANTERIOR,
        COL_FIADO_USADO,
    }
)
EXPORT_OCULTAS = frozenset({COL_ID})
IMPORT_EDIT_KEYS = frozenset(
    {
        COL_NOME,
        COL_WHATSAPP,
        COL_CPF,
        COL_ATIVO,
        COL_CEP,
        COL_UF,
        COL_CIDADE,
        COL_BAIRRO,
        COL_LOGRADOURO,
        COL_NUMERO,
        COL_COMPLEMENTO,
        COL_PLUS_CODE,
        COL_REF_RURAL,
        COL_MAPS_URL,
        COL_SALDO_CASHBACK,
        COL_SALDO_VALE,
        COL_LIMITE_FIADO,
    }
)


def _norm_header(h: str) -> str:
    s = unicodedata.normalize("NFD", str(h or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.lower().strip()
    return re.sub(r"\s+", " ", s)


def _map_headers(headers: list[str]) -> dict[str, str | None]:
    norm = {_norm_header(h): h for h in headers if str(h or "").strip()}
    aliases: dict[str, tuple[str, ...]] = {
        COL_ID: ("id", "codigo agro", "codigo_agro", "pk"),
        COL_CODIGO_ERP: ("codigo erp", "codigo_erp", "externo id", "id erp"),
        COL_NOME: ("nome", "cliente"),
        COL_WHATSAPP: ("whatsapp", "wa", "telefone", "celular"),
        COL_CPF: ("cpf", "documento"),
        COL_ATIVO: ("ativo", "status"),
        COL_CEP: ("cep",),
        COL_UF: ("uf", "estado"),
        COL_CIDADE: ("cidade",),
        COL_BAIRRO: ("bairro",),
        COL_LOGRADOURO: ("logradouro", "rua", "endereco logradouro"),
        COL_NUMERO: ("numero", "número", "n"),
        COL_COMPLEMENTO: ("complemento",),
        COL_PLUS_CODE: ("plus code", "plus_code", "plus code maps"),
        COL_REF_RURAL: ("referencia rural", "referencia_rural", "referencia entrega"),
        COL_MAPS_URL: ("link maps", "maps url", "maps_url_manual", "url maps"),
        COL_ENDERECO: ("endereco resumo", "endereço resumo", "endereco", "endereço"),
        COL_SALDO_CASHBACK: ("saldo cashback", "cashback"),
        COL_SALDO_VALE: ("saldo vale", "vale credito", "vale crédito"),
        COL_LIMITE_FIADO: ("limite fiado", "limite_fiado_local", "limite de fiado"),
        COL_QTD_FIADO_3M: ("qtd compras fiado 3m", "qtd compras fiado (3 meses)"),
        COL_MEDIA_FIADO_3M: ("media fiado 3m", "média fiado (3 meses)", "media fiado (3 meses)"),
        COL_MES_MAIS_FIADO: ("mes mais comprou fiado", "mês que mais comprou fiado"),
        COL_VALOR_MES_MAIS: (
            "valor do mes que mais comprou",
            "valor do mês que mais comprou",
        ),
        COL_VALOR_MES_ANTERIOR: ("valor mes anterior", "valor mês anterior"),
        COL_FIADO_USADO: ("fiado em aberto", "fiado usado agora"),
    }
    out: dict[str, str | None] = {}
    for key, keys in aliases.items():
        out[key] = None
        for k in keys:
            if k in norm:
                out[key] = norm[k]
                break
    return out


def _cel_str(val) -> str:
    if val is None:
        return ""
    if isinstance(val, bool):
        return "Sim" if val else "Não"
    if isinstance(val, int):
        return str(val)
    if isinstance(val, float):
        if abs(val - round(val)) < 1e-9:
            return str(int(round(val)))
        return str(val).strip()
    s = str(val).strip()
    if s.endswith(".0") and s[:-2].replace("-", "").isdigit():
        return s[:-2]
    return s


def _parse_decimal_br(val) -> Decimal | None:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return Decimal(str(val)).quantize(Decimal("0.01"))
    s = str(val).strip()
    if not s or s in ("-", "—"):
        return None
    s = s.replace("R$", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def _parse_bool(val) -> bool | None:
    s = _cel_str(val).lower()
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    if s in ("1", "s", "sim", "true", "yes", "ativo", "on"):
        return True
    if s in ("0", "n", "nao", "false", "no", "inativo", "off"):
        return False
    return None


def _sim_nao(flag: bool) -> str:
    return "Sim" if flag else "Não"


def _ler_planilha(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    suf = path.suffix.lower()
    if suf == ".csv":
        import csv

        for enc in ("utf-8-sig", "latin-1", "cp1252"):
            try:
                with path.open("r", encoding=enc, newline="") as f:
                    reader = csv.DictReader(f, delimiter=";")
                    if reader.fieldnames and len(reader.fieldnames) == 1:
                        f.seek(0)
                        reader = csv.DictReader(f, delimiter=",")
                    return list(reader.fieldnames or []), [dict(r) for r in reader]
            except UnicodeDecodeError:
                continue
        raise ValueError("Não foi possível ler o CSV.")
    if suf in (".xlsx", ".xls"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(c or "").strip() for c in next(it, [])]
        rows = []
        for row in it:
            if not any(row):
                continue
            d = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers)) if headers[i]}
            rows.append(d)
        wb.close()
        return headers, rows
    raise ValueError("Use arquivo .csv ou .xlsx.")


def _label_mes(ano: int, mes: int) -> str:
    if 1 <= mes <= 12:
        return f"{MESES_PT[mes - 1]}/{ano}"
    return f"{mes:02d}/{ano}"


def _data_local(dt) -> date:
    """Dia da loja (Jacupiranga), não o dia em UTC."""
    if dt is None:
        return timezone.localdate()
    if timezone.is_aware(dt):
        return timezone.localtime(dt).date()
    return dt.date()


def _fiado_stats_3m_por_cliente() -> dict[int, dict[str, Any]]:
    """Compras fiado por venda nos últimos ~3 meses (90 dias)."""
    hoje = timezone.localdate()
    desde = hoje - timedelta(days=90)
    por_cliente: dict[int, list[tuple[date, Decimal]]] = defaultdict(list)

    vendas = VendaAgro.objects.filter(
        devolvida_em__isnull=True,
        criado_em__date__gte=desde,
    ).only("pk", "cliente_id_erp", "pagamentos_json", "forma_pagamento", "total", "criado_em")
    for venda in vendas.iterator(chunk_size=400):
        valor = valor_fiado_venda_local(venda)
        if valor <= Decimal("0.009"):
            continue
        _erp, agro_pk, _cli = resolver_cliente_fiado(venda.cliente_id_erp)
        if not agro_pk:
            agro_pk = cliente_agro_pk_de_ref(venda.cliente_id_erp)
        if not agro_pk or not venda.criado_em:
            continue
        por_cliente[int(agro_pk)].append((_data_local(venda.criado_em), valor.quantize(Decimal("0.01"))))

    titulos = (
        FiadoTituloAgro.objects.filter(venda_agro_id__isnull=True, criado_em__date__gte=desde)
        .exclude(situacao=FiadoTituloAgro.Situacao.CANCELADO)
        .only("cliente_agro_id", "numero_documento", "valor_bruto", "chave_unica", "criado_em")
    )
    agrupado: dict[tuple[int, str], tuple[date, Decimal]] = {}
    for t in titulos.iterator(chunk_size=400):
        if not t.cliente_agro_id:
            continue
        doc = (t.numero_documento or t.chave_unica or f"t{t.pk}").strip()
        key = (int(t.cliente_agro_id), doc)
        dt = _data_local(t.criado_em) if t.criado_em else hoje
        val = _parse_decimal_br(t.valor_bruto) or Decimal("0")
        if key in agrupado:
            prev_dt, prev_val = agrupado[key]
            agrupado[key] = (min(prev_dt, dt), (prev_val + val).quantize(Decimal("0.01")))
        else:
            agrupado[key] = (dt, val.quantize(Decimal("0.01")))
    for (pk, _doc), (dt, val) in agrupado.items():
        if val > Decimal("0.009"):
            por_cliente[pk].append((dt, val))

    out: dict[int, dict[str, Any]] = {}
    for pk, compras in por_cliente.items():
        qtd = len(compras)
        total = sum(v for _d, v in compras).quantize(Decimal("0.01"))
        media = (total / qtd).quantize(Decimal("0.01")) if qtd else Decimal("0")
        cont_mes = Counter((d.year, d.month) for d, _v in compras)
        soma_mes: dict[tuple[int, int], Decimal] = defaultdict(lambda: Decimal("0"))
        for d, v in compras:
            soma_mes[(d.year, d.month)] += v
        mes_top = ""
        valor_top = Decimal("0")
        if cont_mes:
            (ano, mes), _q = cont_mes.most_common(1)[0]
            mes_top = _label_mes(ano, mes)
            valor_top = soma_mes[(ano, mes)].quantize(Decimal("0.01"))
        ant = hoje.replace(day=1) - timedelta(days=1)
        valor_ant = soma_mes.get((ant.year, ant.month), Decimal("0")).quantize(Decimal("0.01"))
        out[pk] = {
            "qtd": qtd,
            "media": float(media),
            "mes_mais": mes_top,
            "valor_mes_mais": float(valor_top),
            "valor_mes_anterior": float(valor_ant),
        }
    return out


def _fiado_usado_mapa() -> dict[int, float]:
    from django.db.models import F, Sum

    qs = (
        FiadoTituloAgro.objects.filter(cliente_agro_id__isnull=False)
        .exclude(
            situacao__in=(
                FiadoTituloAgro.Situacao.QUITADO,
                FiadoTituloAgro.Situacao.CANCELADO,
            )
        )
        .values("cliente_agro_id")
        .annotate(usado=Sum(F("valor_bruto") - F("valor_pago")))
    )
    out: dict[int, float] = {}
    for row in qs:
        pk = int(row["cliente_agro_id"])
        val = _parse_decimal_br(row.get("usado")) or Decimal("0")
        if val > Decimal("0.009"):
            out[pk] = float(val.quantize(Decimal("0.01")))
    return out


def coletar_linhas_export_clientes() -> list[dict[str, Any]]:
    stats = _fiado_stats_3m_por_cliente()
    usado_map = _fiado_usado_mapa()
    linhas: list[dict[str, Any]] = []
    qs = ClienteAgro.objects.all().order_by("nome")[:EXPORT_MAX_ROWS]
    for cli in qs.iterator(chunk_size=300):
        st = stats.get(cli.pk) or {}
        linhas.append(
            {
                COL_ID: cli.pk,
                COL_CODIGO_ERP: (cli.externo_id or "").strip(),
                COL_NOME: cli.nome,
                COL_WHATSAPP: cli.whatsapp,
                COL_CPF: cli.cpf,
                COL_ATIVO: _sim_nao(cli.ativo),
                COL_CEP: cli.cep,
                COL_UF: cli.uf,
                COL_CIDADE: cli.cidade,
                COL_BAIRRO: cli.bairro,
                COL_LOGRADOURO: cli.logradouro,
                COL_NUMERO: cli.numero,
                COL_COMPLEMENTO: cli.complemento,
                COL_PLUS_CODE: cli.plus_code,
                COL_REF_RURAL: cli.referencia_rural,
                COL_MAPS_URL: cli.maps_url_manual,
                COL_ENDERECO: cli.endereco,
                COL_SALDO_CASHBACK: float(cli.saldo_cashback or 0),
                COL_SALDO_VALE: float(cli.saldo_vale_credito or 0),
                COL_LIMITE_FIADO: float(cli.limite_fiado_local or 0),
                COL_QTD_FIADO_3M: int(st.get("qtd") or 0),
                COL_MEDIA_FIADO_3M: float(st.get("media") or 0),
                COL_MES_MAIS_FIADO: st.get("mes_mais") or "",
                COL_VALOR_MES_MAIS: float(st.get("valor_mes_mais") or 0),
                COL_VALOR_MES_ANTERIOR: float(st.get("valor_mes_anterior") or 0),
                COL_FIADO_USADO: float(usado_map.get(cli.pk) or 0),
            }
        )
    return linhas


def montar_xlsx_clientes(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    fill_edit = PatternFill("solid", fgColor="FFF9C4")
    fill_ro = PatternFill("solid", fgColor="F1F5F9")

    for col, (label, key) in enumerate(EXPORT_HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=label)
        c.font = Font(bold=True)
        if key in EXPORT_ONLY:
            c.fill = fill_ro
        else:
            c.fill = fill_edit

    for r_idx, row in enumerate(rows, start=2):
        for col, (_label, key) in enumerate(EXPORT_HEADERS, start=1):
            val = row.get(key, "")
            cell = ws.cell(row=r_idx, column=col, value=val)
            if key in EXPORT_OCULTAS:
                cell.protection = Protection(locked=True, hidden=True)
            elif key in EXPORT_ONLY:
                cell.protection = Protection(locked=True)
                cell.fill = fill_ro
            else:
                cell.protection = Protection(locked=False)
                cell.fill = fill_edit

    for col, (label, key) in enumerate(EXPORT_HEADERS, start=1):
        letter = get_column_letter(col)
        if key in EXPORT_OCULTAS:
            ws.column_dimensions[letter].hidden = True
            ws.column_dimensions[letter].width = 2
        elif key == COL_NOME:
            ws.column_dimensions[letter].width = 36
        elif key == COL_ENDERECO:
            ws.column_dimensions[letter].width = 42
        else:
            ws.column_dimensions[letter].width = max(14, min(28, len(label) + 2))

    ws2 = wb.create_sheet("Como usar")
    ws2["A1"] = "Clientes — Excel"
    ws2["A1"].font = Font(bold=True, size=14)
    dicas = [
        "",
        "1) Excel ↓ baixa todos os clientes do sistema (dados online da loja).",
        "2) Colunas amarelas podem ser editadas. Cinza = só leitura (média fiado, mês que mais comprou, etc.).",
        "3) Média fiado (3 meses) = valor médio de cada COMPRA fiado nos últimos 90 dias (não é por parcela).",
        "4) Mês que mais comprou fiado = mês com mais compras (ex.: Ago/2026). A coluna ao lado é o valor comprado nesse mês.",
        "4b) Valor mês anterior = quanto comprou fiado no mês calendário anterior (só leitura).",
        "5) Limite fiado: 0 = volta ao padrão (R$ 5.000). 0,01 = bloqueia fiado. Qualquer outro valor = limite fixo.",
        "6) Não apague a coluna ID (oculta). Linha sem ID não é importada.",
        "7) Célula vazia na importação = não muda aquele campo.",
        "8) Excel ↑ mostra prévia antes de gravar — confira e confirme.",
    ]
    for i, t in enumerate(dicas, start=2):
        ws2.cell(row=i, column=1, value=t)
    ws2.column_dimensions["A"].width = 105

    ws.protection.sheet = True
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _json_safe(val: Any) -> Any:
    if isinstance(val, Decimal):
        return float(val.quantize(Decimal("0.01")))
    return val


def _cliente_para_dict(cli: ClienteAgro) -> dict[str, Any]:
    return {
        COL_NOME: cli.nome,
        COL_WHATSAPP: cli.whatsapp,
        COL_CPF: cli.cpf,
        COL_ATIVO: cli.ativo,
        COL_CEP: cli.cep,
        COL_UF: cli.uf,
        COL_CIDADE: cli.cidade,
        COL_BAIRRO: cli.bairro,
        COL_LOGRADOURO: cli.logradouro,
        COL_NUMERO: cli.numero,
        COL_COMPLEMENTO: cli.complemento,
        COL_PLUS_CODE: cli.plus_code,
        COL_REF_RURAL: cli.referencia_rural,
        COL_MAPS_URL: cli.maps_url_manual,
        COL_SALDO_CASHBACK: _json_safe(cli.saldo_cashback),
        COL_SALDO_VALE: _json_safe(cli.saldo_vale_credito),
        COL_LIMITE_FIADO: _json_safe(cli.limite_fiado_local),
    }


def _patch_da_linha(raw: dict, colmap: dict[str, str | None]) -> dict[str, Any]:
    patch: dict[str, Any] = {}

    def txt(key: str, max_len: int) -> None:
        hdr = colmap.get(key)
        if not hdr or hdr not in raw:
            return
        v = _cel_str(raw.get(hdr))
        if not v:
            return
        patch[key] = v[:max_len]

    def dec(key: str) -> None:
        hdr = colmap.get(key)
        if not hdr or hdr not in raw:
            return
        v = raw.get(hdr)
        if v is None or (isinstance(v, str) and not str(v).strip()):
            return
        d = _parse_decimal_br(v)
        if d is None:
            patch[f"__erro_{key}"] = f"Valor inválido em «{hdr}»."
        else:
            patch[key] = d

    txt(COL_NOME, 200)
    txt(COL_WHATSAPP, 20)
    txt(COL_CPF, 14)
    txt(COL_CEP, 12)
    txt(COL_UF, 2)
    txt(COL_CIDADE, 120)
    txt(COL_BAIRRO, 120)
    txt(COL_LOGRADOURO, 300)
    txt(COL_NUMERO, 30)
    txt(COL_COMPLEMENTO, 200)
    txt(COL_PLUS_CODE, 120)
    txt(COL_REF_RURAL, 300)
    txt(COL_MAPS_URL, 600)
    dec(COL_SALDO_CASHBACK)
    dec(COL_SALDO_VALE)
    dec(COL_LIMITE_FIADO)

    hdr_ativo = colmap.get(COL_ATIVO)
    if hdr_ativo and hdr_ativo in raw:
        v = raw.get(hdr_ativo)
        if v is not None and str(v).strip():
            b = _parse_bool(v)
            if b is None:
                patch["__erro_ativo"] = "Ativo: use Sim ou Não."
            else:
                patch[COL_ATIVO] = b
    return patch


def preview_importacao_clientes(path: Path) -> dict[str, Any]:
    headers, rows_raw = _ler_planilha(path)
    colmap = _map_headers(headers)
    hdr_id = colmap.get(COL_ID)
    if not hdr_id:
        raise ValueError("Coluna «ID» não encontrada. Baixe de novo com Excel ↓.")

    alteracoes: list[dict] = []
    ignoradas: list[dict] = []
    erros: list[dict] = []
    vistos: set[int] = set()

    for i, raw in enumerate(rows_raw[:IMPORT_MAX_ROWS], start=2):
        pid_raw = _cel_str(raw.get(hdr_id or ""))
        if not pid_raw:
            continue
        try:
            pk = int(float(pid_raw))
        except (TypeError, ValueError):
            erros.append({"linha": i, "id": pid_raw, "erro": "ID inválido."})
            continue
        if pk in vistos:
            erros.append({"linha": i, "id": pk, "erro": "ID duplicado na planilha."})
            continue
        vistos.add(pk)

        patch = _patch_da_linha(raw, colmap)
        err_fields = [v for k, v in patch.items() if str(k).startswith("__erro_")]
        if err_fields:
            erros.append({"linha": i, "id": pk, "erro": err_fields[0]})
            continue
        edit = {k: v for k, v in patch.items() if k in IMPORT_EDIT_KEYS}
        if not edit:
            ignoradas.append({"linha": i, "id": pk, "motivo": "Nenhum campo editável preenchido."})
            continue

        cli = ClienteAgro.objects.filter(pk=pk).first()
        if not cli:
            erros.append({"linha": i, "id": pk, "erro": "Cliente não encontrado."})
            continue

        atual = _cliente_para_dict(cli)
        campos = []
        for k, novo in edit.items():
            ant = atual.get(k)
            if k in (COL_SALDO_CASHBACK, COL_SALDO_VALE, COL_LIMITE_FIADO):
                a = _parse_decimal_br(ant) or Decimal("0")
                n = _parse_decimal_br(novo) or Decimal("0")
                if a == n:
                    continue
                campos.append({"campo": k, "de": float(a), "para": float(n)})
            elif k == COL_ATIVO:
                if bool(ant) == bool(novo):
                    continue
                campos.append({"campo": k, "de": _sim_nao(bool(ant)), "para": _sim_nao(bool(novo))})
            else:
                if str(ant or "").strip() == str(novo or "").strip():
                    continue
                campos.append({"campo": k, "de": str(ant or ""), "para": str(novo or "")})
        if not campos:
            ignoradas.append({"linha": i, "id": pk, "motivo": "Valores iguais ao cadastro."})
            continue
        alteracoes.append(
            {
                "linha": i,
                "id": pk,
                "nome": cli.nome,
                "campos": campos,
            }
        )

    return {
        "total_linhas": len(rows_raw),
        "alteracoes": alteracoes[:400],
        "n_alteracoes": len(alteracoes),
        "ignoradas": ignoradas[:80],
        "n_ignoradas": len(ignoradas),
        "erros": erros[:120],
        "n_erros": len(erros),
    }


def _aplicar_patch_cliente(cli: ClienteAgro, patch: dict[str, Any], user) -> list[str]:
    from produtos.fiado_gestao_util import definir_limite_fiado_cliente
    from produtos.models import compor_endereco_resumo_cliente

    nome_antes = cli.nome

    alterados: list[str] = []
    endereco_keys = {
        COL_CEP,
        COL_UF,
        COL_CIDADE,
        COL_BAIRRO,
        COL_LOGRADOURO,
        COL_NUMERO,
        COL_COMPLEMENTO,
    }

    if COL_NOME in patch:
        cli.nome = str(patch[COL_NOME])[:200]
        alterados.append(COL_NOME)
    if COL_WHATSAPP in patch:
        cli.whatsapp = str(patch[COL_WHATSAPP])[:20]
        alterados.append(COL_WHATSAPP)
    if COL_CPF in patch:
        cli.cpf = str(patch[COL_CPF])[:14]
        alterados.append(COL_CPF)
    if COL_ATIVO in patch:
        cli.ativo = bool(patch[COL_ATIVO])
        alterados.append(COL_ATIVO)
    for k in endereco_keys:
        if k in patch:
            max_len = ClienteAgro._meta.get_field(k).max_length or 120
            setattr(cli, k, str(patch[k])[:max_len])
            alterados.append(k)
    if COL_PLUS_CODE in patch:
        cli.plus_code = str(patch[COL_PLUS_CODE])[:120]
        alterados.append(COL_PLUS_CODE)
    if COL_REF_RURAL in patch:
        cli.referencia_rural = str(patch[COL_REF_RURAL])[:300]
        alterados.append(COL_REF_RURAL)
    if COL_MAPS_URL in patch:
        cli.maps_url_manual = str(patch[COL_MAPS_URL])[:600]
        alterados.append(COL_MAPS_URL)
    if any(k in patch for k in endereco_keys):
        cli.endereco = compor_endereco_resumo_cliente(
            cli.cep,
            cli.uf,
            cli.cidade,
            cli.bairro,
            cli.logradouro,
            cli.numero,
            cli.complemento,
        )[:500]
        alterados.append(COL_ENDERECO)
    if COL_SALDO_CASHBACK in patch:
        cli.saldo_cashback = patch[COL_SALDO_CASHBACK]
        alterados.append(COL_SALDO_CASHBACK)
    if COL_SALDO_VALE in patch:
        cli.saldo_vale_credito = patch[COL_SALDO_VALE]
        alterados.append(COL_SALDO_VALE)

    limite_patch = patch.get(COL_LIMITE_FIADO)
    if COL_LIMITE_FIADO in patch:
        alterados.append(COL_LIMITE_FIADO)

    if alterados:
        cli.editado_local = True
        fields = [f for f in alterados if f != COL_LIMITE_FIADO]
        fields.append("editado_local")
        fields.append("atualizado_em")
        cli.save(update_fields=list(dict.fromkeys(fields)))

    if COL_LIMITE_FIADO in patch:
        usuario = ""
        if user and getattr(user, "is_authenticated", False):
            usuario = user.get_username() or str(user.pk)
        definir_limite_fiado_cliente(
            cli.pk,
            limite_patch,
            usuario=usuario or "planilha-clientes",
        )
    if COL_NOME in alterados:
        from produtos.cliente_operacoes_util import propagar_renome_cliente

        propagar_renome_cliente(cli, nome_antes)
    return alterados


def aplicar_importacao_clientes(path: Path, user, *, nome_arquivo: str = "") -> dict[str, Any]:
    prev = preview_importacao_clientes(path)
    if not prev.get("n_alteracoes"):
        raise ValueError("Nenhuma alteração para gravar.")

    backups: list[dict] = []
    ok = 0
    with transaction.atomic():
        for item in prev["alteracoes"]:
            pk = int(item["id"])
            cli = ClienteAgro.objects.filter(pk=pk).first()
            if not cli:
                continue
            snap_antes = _cliente_para_dict(cli)
            patch = {c["campo"]: c["para"] for c in item["campos"]}
            alterados = _aplicar_patch_cliente(cli, patch, user)
            if alterados:
                backups.append(
                    {
                        "id": pk,
                        "nome": cli.nome,
                        "antes": {k: _json_safe(snap_antes.get(k)) for k in alterados},
                        "para": {k: _json_safe(patch.get(k)) for k in alterados},
                        "campos_alterados": alterados,
                    }
                )
                ok += 1

        hist = CadastroPlanilhaImportHistoricoAgro.objects.create(
            usuario=user if user and getattr(user, "is_authenticated", False) else None,
            nome_arquivo=(nome_arquivo or path.name)[:255],
            n_produtos=ok,
            n_campos=sum(len(b.get("campos_alterados") or []) for b in backups),
            tipo=CadastroPlanilhaImportHistoricoAgro.Tipo.CADASTRO,
            backup={"tipo": TIPO_HISTORICO, "items": backups},
        )

    return {
        "historico_id": hist.pk,
        "clientes_alterados": ok,
        "n_campos": hist.n_campos,
        "preview": {
            "n_erros": prev.get("n_erros"),
            "n_ignoradas": prev.get("n_ignoradas"),
        },
    }
