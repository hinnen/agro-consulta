"""Venda hora a hora — Central de Relatórios.

Fonte: VendaAgro (devolução total fica de fora). Hora no fuso da loja.
Expediente exibido: 7h–18h (a loja abre 7h30 e fecha 18h30).
Costume: mesma hora, mesmos dias da semana, nas 4 semanas anteriores ao período.
"""
from __future__ import annotations

import logging
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from produtos.relatorios_vendas_util import fmt_brl, xlsx_http_response

logger = logging.getLogger(__name__)

HORA_INI = 7
HORA_FIM = 18
DIAS_SEMANA = ("Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom")
PERIODO_MAX_DIAS = 92
DETALHE_LIMITE = 100


def _aware(d0: date, d1: date) -> tuple[datetime, datetime]:
    tz = timezone.get_current_timezone()
    desde = timezone.make_aware(datetime.combine(d0, time.min), tz)
    ate = timezone.make_aware(datetime.combine(d1, time.max), tz)
    return desde, ate


def _parse_d(raw: str) -> date | None:
    try:
        return date.fromisoformat((raw or "").strip()[:10])
    except (TypeError, ValueError):
        return None


def listar_dias(d0: date, d1: date) -> list[date]:
    out: list[date] = []
    d = d0
    while d <= d1:
        out.append(d)
        d += timedelta(days=1)
    return out


def dias_costume(dias: list[date]) -> list[date]:
    """Até 4 ocorrências de cada dia da semana do período, nas 4 semanas anteriores."""
    if not dias:
        return []
    inicio = min(dias)
    quer = {d.weekday() for d in dias}
    out: list[date] = []
    for i in range(1, 29):
        d = inicio - timedelta(days=i)
        if d.weekday() in quer:
            out.append(d)
    out.sort()
    return out


def parse_periodo_hora(request) -> dict[str, Any]:
    """hoje | ontem | semana | 7d | 30d | mes | custom. Padrão: hoje. Teto: 92 dias."""
    hoje = timezone.localdate()
    periodo = (request.GET.get("periodo") or "hoje").strip().lower()
    d_de = _parse_d(request.GET.get("de") or "")
    d_ate = _parse_d(request.GET.get("ate") or "")
    aviso = ""

    if periodo == "ontem":
        d0 = d1 = hoje - timedelta(days=1)
    elif periodo == "semana":
        d0 = hoje - timedelta(days=hoje.weekday())
        d1 = hoje
    elif periodo == "7d":
        d0, d1 = hoje - timedelta(days=6), hoje
    elif periodo == "30d":
        d0, d1 = hoje - timedelta(days=29), hoje
    elif periodo in ("mes", "mes_atual"):
        periodo = "mes"
        d0, d1 = hoje.replace(day=1), hoje
    elif periodo == "custom":
        d0 = d_de or hoje
        d1 = d_ate or hoje
    else:
        periodo = "hoje"
        d0 = d1 = hoje

    if d_de and d_ate and periodo != "custom" and (d_de != d0 or d_ate != d1):
        periodo = "custom"
        d0, d1 = d_de, d_ate
    if d0 > d1:
        d0, d1 = d1, d0
    if (d1 - d0).days + 1 > PERIODO_MAX_DIAS:
        d0 = d1 - timedelta(days=PERIODO_MAX_DIAS - 1)
        periodo = "custom"
        aviso = f"Período limitado aos últimos {PERIODO_MAX_DIAS} dias."

    return {
        "periodo": periodo,
        "de": d0.isoformat(),
        "ate": d1.isoformat(),
        "d0": d0,
        "d1": d1,
        "label": f"{d0.strftime('%d/%m/%Y')} — {d1.strftime('%d/%m/%Y')}",
        "aviso": aviso,
    }


def parse_loja_hora(request) -> str:
    raw = (request.GET.get("deposito") or "ambos").strip().lower()
    if raw in ("centro", "vila"):
        return raw
    return "ambos"


def parse_canal_hora(request) -> str:
    raw = (request.GET.get("canal") or "todos").strip().lower()
    if raw in ("balcao", "entrega"):
        return raw
    return "todos"


def parse_visao_hora(request, n_dias: int) -> str:
    if n_dias <= 1:
        return "soma"
    raw = (request.GET.get("visao") or "media").strip().lower()
    if raw == "soma":
        return "soma"
    return "media"


def parse_hora_detalhe(request) -> int | None:
    raw = (request.GET.get("hora") or "").strip()
    if not raw.isdigit():
        return None
    h = int(raw)
    if 0 <= h <= 23:
        return h
    return None


def parse_wd_detalhe(request) -> int | None:
    raw = (request.GET.get("wd") or "").strip()
    if not raw.isdigit():
        return None
    wd = int(raw)
    if 0 <= wd <= 6:
        return wd
    return None


def _como_local(dt: datetime | None) -> datetime:
    if dt is None:
        return datetime.min
    if timezone.is_aware(dt):
        dt = timezone.localtime(dt)
    return dt.replace(tzinfo=None)


def _loja_de(dep: str | None) -> str:
    return "vila" if str(dep or "").strip().lower() == "vila" else "centro"


def _passa(venda: dict, loja: str, canal: str) -> bool:
    if loja in ("centro", "vila") and _loja_de(venda.get("deposito")) != loja:
        return False
    entrega = bool(venda.get("entrega"))
    if canal == "balcao" and entrega:
        return False
    if canal == "entrega" and not entrega:
        return False
    return True


def _fmt_n(n: float, media: bool) -> str:
    if not media:
        return str(int(round(n)))
    if abs(n - round(n)) < 0.05:
        return str(int(round(n)))
    return f"{n:.1f}".replace(".", ",")


def _fmt_pct(v: float | None) -> str:
    if v is None:
        return "—"
    s = f"{v:.1f}".replace(".", ",")
    return s + "%"


def _var_label(var: float | None, base: float | None) -> str:
    if base is None or abs(base) < 0.005:
        if var is None:
            return "sem costume"
        return "igual"
    if abs(var) < 0.05:
        return "igual"
    txt = f"{var:+.1f}".replace(".", ",")
    return txt + "%"


def _novo_bucket() -> dict:
    return {
        "centro": 0.0,
        "vila": 0.0,
        "n_centro": 0,
        "n_vila": 0,
        "balcao": 0.0,
        "entrega": 0.0,
        "ops": Counter(),
    }


def _soma_venda(bucket: dict, venda: dict, valor: float) -> None:
    loja = _loja_de(venda.get("deposito"))
    if loja == "vila":
        bucket["vila"] += valor
        bucket["n_vila"] += 1
    else:
        bucket["centro"] += valor
        bucket["n_centro"] += 1
    if venda.get("entrega"):
        bucket["entrega"] += valor
    else:
        bucket["balcao"] += valor
    op = (venda.get("operador") or "").strip() or "Sem nome"
    bucket["ops"][op] += 1


def _quem(ops: Counter, limite: int = 3) -> str:
    if not ops:
        return "—"
    pares = sorted(ops.items(), key=lambda kv: (-kv[1], kv[0].casefold()))
    partes = []
    for nome, n in pares[:limite]:
        curto = nome if len(nome) <= 22 else nome[:21] + "…"
        partes.append(f"{curto} {n}")
    extra = len(pares) - limite
    if extra > 0:
        partes.append(f"+{extra}")
    return " · ".join(partes)


def agregar_hora(
    vendas: list[dict],
    baseline: list[dict],
    *,
    dias: list[date],
    dias_base: list[date],
    loja: str = "ambos",
    canal: str = "todos",
    visao: str = "soma",
    hora_detalhe: int | None = None,
    wd_detalhe: int | None = None,
) -> dict[str, Any]:
    """Monta a grade, o mapa e o detalhe. Não consulta banco."""
    n_dias = max(len(dias), 1)
    media = visao == "media" and n_dias > 1
    divisor = float(n_dias) if media else 1.0
    base_set = set(dias_base)
    n_base = len(dias_base)

    por_dia: dict[tuple[date, int], dict] = defaultdict(_novo_bucket)
    base_hora = [_novo_bucket() for _ in range(24)]
    lista_ok: list[dict] = []
    dia_set = set(dias)

    for venda in vendas or []:
        if not _passa(venda, loja, canal):
            continue
        quando = _como_local(venda.get("quando"))
        if dia_set and quando.date() not in dia_set:
            continue
        valor = float(venda.get("total") or 0)
        hora = quando.hour
        _soma_venda(por_dia[(quando.date(), hora)], venda, valor)
        lista_ok.append({**venda, "_quando": quando, "_valor": valor})

    for venda in baseline or []:
        if not _passa(venda, loja, canal):
            continue
        quando = _como_local(venda.get("quando"))
        if base_set and quando.date() not in base_set:
            continue
        _soma_venda(base_hora[quando.hour], venda, float(venda.get("total") or 0))

    def _total_bucket(b: dict) -> float:
        return b["centro"] + b["vila"]

    def _n_bucket(b: dict) -> int:
        return b["n_centro"] + b["n_vila"]

    horas_sum = [_novo_bucket() for _ in range(24)]
    for (dia, hora), b in por_dia.items():
        dest = horas_sum[hora]
        for k in ("centro", "vila", "balcao", "entrega"):
            dest[k] += b[k]
        dest["n_centro"] += b["n_centro"]
        dest["n_vila"] += b["n_vila"]
        dest["ops"].update(b["ops"])

    grand = sum(_total_bucket(b) for b in horas_sum)
    grand_n = sum(_n_bucket(b) for b in horas_sum)

    def _cmp_base(hora: int) -> float | None:
        if n_base <= 0:
            return None
        media_base = _total_bucket(base_hora[hora]) / float(n_base)
        if media:
            return media_base
        return media_base * float(n_dias)

    linhas = []
    for hora in range(24):
        b = horas_sum[hora]
        total = _total_bucket(b)
        n = _n_bucket(b)
        atual = total / divisor
        base = _cmp_base(hora)
        var = None
        if base is not None and abs(base) >= 0.005:
            var = round((atual - base) / base * 100.0, 1)
        elif base is not None and abs(atual) < 0.005:
            var = 0.0
        linhas.append(
            {
                "hora": hora,
                "rotulo": f"{hora}h",
                "expediente": HORA_INI <= hora <= HORA_FIM,
                "centro": round(b["centro"] / divisor, 2),
                "vila": round(b["vila"] / divisor, 2),
                "total": round(atual, 2),
                "n": n / divisor,
                "n_centro": b["n_centro"] / divisor,
                "n_vila": b["n_vila"] / divisor,
                "ticket": round(total / n, 2) if n else 0.0,
                "pct": round(total / grand * 100.0, 1) if grand else 0.0,
                "balcao": round(b["balcao"] / divisor, 2),
                "entrega": round(b["entrega"] / divisor, 2),
                "base": None if base is None else round(base, 2),
                "var": var,
                "quem": _quem(b["ops"]),
                "total_bruto": round(total, 2),
                "n_bruto": n,
            }
        )

    pico_max = max((ln["total_bruto"] for ln in linhas if ln["expediente"]), default=0.0)
    candidatos_pico = [ln for ln in linhas if ln["expediente"] and ln["total_bruto"] == pico_max and pico_max > 0]
    pico = candidatos_pico[0] if candidatos_pico else None
    com_venda = [ln for ln in linhas if ln["expediente"] and ln["total_bruto"] > 0]
    fraca = None
    if len(com_venda) >= 2:
        menor = min(ln["total_bruto"] for ln in com_venda)
        fraca = next(ln for ln in com_venda if ln["total_bruto"] == menor)
        if pico and fraca["hora"] == pico["hora"]:
            fraca = None

    escala = max((ln["total"] for ln in linhas), default=0.0)
    for ln in linhas:
        ln["barra"] = int(round(ln["total"] / escala * 100)) if escala > 0 else 0
        ln["pico"] = bool(pico and ln["hora"] == pico["hora"])
        ln["fraca"] = bool(fraca and ln["hora"] == fraca["hora"])
        ln["total_fmt"] = fmt_brl(ln["total"])
        ln["centro_fmt"] = fmt_brl(ln["centro"])
        ln["vila_fmt"] = fmt_brl(ln["vila"])
        ln["ticket_fmt"] = fmt_brl(ln["ticket"]) if ln["n_bruto"] else "—"
        ln["balcao_fmt"] = fmt_brl(ln["balcao"])
        ln["entrega_fmt"] = fmt_brl(ln["entrega"])
        ln["base_fmt"] = fmt_brl(ln["base"]) if ln["base"] is not None else "—"
        ln["pct_fmt"] = _fmt_pct(ln["pct"])
        ln["var_label"] = _var_label(ln["var"], ln["base"])
        if ln["var"] is None or abs(ln["var"]) < 0.05:
            ln["var_tom"] = "neutro"
        elif ln["var"] > 0:
            ln["var_tom"] = "acima"
        else:
            ln["var_tom"] = "abaixo"
        ln["n_fmt"] = _fmt_n(ln["n"], media)
        ln["n_centro_fmt"] = _fmt_n(ln["n_centro"], media)
        ln["n_vila_fmt"] = _fmt_n(ln["n_vila"], media)

    # Mapa: média daquela hora naquele dia da semana, dentro do período.
    mapa = []
    mapa_max = 0.0
    dias_por_wd: dict[int, list[date]] = defaultdict(list)
    for d in dias:
        dias_por_wd[d.weekday()].append(d)
    for wd in range(7):
        ocorrencias = dias_por_wd.get(wd) or []
        celulas = []
        for hora in range(HORA_INI, HORA_FIM + 1):
            if not ocorrencias:
                celulas.append({"hora": hora, "valor": None, "vazio": True})
                continue
            soma = 0.0
            for d in ocorrencias:
                soma += _total_bucket(por_dia.get((d, hora)) or _novo_bucket())
            valor = round(soma / float(len(ocorrencias)), 2)
            mapa_max = max(mapa_max, valor)
            celulas.append({"hora": hora, "valor": valor, "vazio": valor <= 0})
        mapa.append({"wd": wd, "rotulo": DIAS_SEMANA[wd], "celulas": celulas, "tem": bool(ocorrencias)})
    for linha in mapa:
        for cel in linha["celulas"]:
            if cel["vazio"] or not cel["valor"] or mapa_max <= 0:
                cel["alpha"] = 0
                cel["valor_fmt"] = "—"
                cel["valor_curto"] = "·"
            else:
                cel["alpha"] = round(0.18 + 0.82 * (cel["valor"] / mapa_max), 2)
                cel["valor_fmt"] = fmt_brl(cel["valor"])
                if cel["valor"] >= 1000:
                    cel["valor_curto"] = f"{cel['valor'] / 1000:.1f}".replace(".", ",") + " mil"
                else:
                    cel["valor_curto"] = f"{cel['valor']:.0f}"

    detalhe = None
    if hora_detalhe is not None:
        escolhidas = []
        for venda in lista_ok:
            quando = venda["_quando"]
            if quando.hour != hora_detalhe:
                continue
            if wd_detalhe is not None and quando.weekday() != wd_detalhe:
                continue
            escolhidas.append(venda)
        escolhidas.sort(key=lambda v: v["_quando"], reverse=True)
        n_det = len(escolhidas)
        visiveis = escolhidas[:DETALHE_LIMITE]
        rotulo = f"{hora_detalhe}h"
        if wd_detalhe is not None:
            rotulo = f"{DIAS_SEMANA[wd_detalhe]} · {rotulo}"
        detalhe = {
            "hora": hora_detalhe,
            "wd": wd_detalhe,
            "rotulo": rotulo,
            "n": n_det,
            "mostrando": len(visiveis),
            "cortou": n_det > len(visiveis),
            "vendas": [
                {
                    "id": v.get("id"),
                    "quando_fmt": v["_quando"].strftime("%d/%m %H:%M"),
                    "loja": "Vila" if _loja_de(v.get("deposito")) == "vila" else "Centro",
                    "cliente": (v.get("cliente") or "").strip() or "Sem cliente",
                    "operador": (v.get("operador") or "").strip() or "Sem nome",
                    "canal": "Entrega" if v.get("entrega") else "Balcão",
                    "total": round(float(v["_valor"]), 2),
                    "total_fmt": fmt_brl(v["_valor"]),
                }
                for v in visiveis
            ],
        }

    total_centro = sum(b["centro"] for b in horas_sum)
    total_vila = sum(b["vila"] for b in horas_sum)
    total_balcao = sum(b["balcao"] for b in horas_sum)
    total_entrega = sum(b["entrega"] for b in horas_sum)
    if media:
        aviso_visao = (
            f"Números em média por dia ({n_dias} dias). "
            "O costume é a média da mesma hora, no mesmo dia da semana, nas 4 semanas anteriores."
        )
    elif n_dias == 1:
        aviso_visao = (
            "Um dia: valor vendido na hora. "
            "O costume compara com a média dessa mesma hora nas 4 semanas anteriores (mesmo dia da semana)."
        )
    else:
        aviso_visao = (
            f"Soma dos {n_dias} dias. "
            "O costume é essa média antiga multiplicada pela quantidade de dias, para a comparação ficar justa."
        )

    def _card(ln: dict | None) -> dict | None:
        if not ln:
            return None
        return {
            "hora": ln["hora"],
            "rotulo": ln["rotulo"],
            "total_fmt": ln["total_fmt"],
            "n_fmt": ln["n_fmt"],
        }

    return {
        "n_dias": n_dias,
        "n_base": n_base,
        "visao": "media" if media else "soma",
        "media": media,
        "loja": loja,
        "duas_lojas": loja == "ambos",
        "canal": canal,
        "aviso_visao": aviso_visao,
        "total": round(grand / divisor, 2),
        "total_bruto": round(grand, 2),
        "total_fmt": fmt_brl(grand / divisor),
        "n_vendas": grand_n / divisor,
        "n_vendas_fmt": _fmt_n(grand_n / divisor, media),
        "ticket": round(grand / grand_n, 2) if grand_n else 0.0,
        "ticket_fmt": fmt_brl(grand / grand_n) if grand_n else "—",
        "centro_fmt": fmt_brl(total_centro / divisor),
        "vila_fmt": fmt_brl(total_vila / divisor),
        "balcao_fmt": fmt_brl(total_balcao / divisor),
        "entrega_fmt": fmt_brl(total_entrega / divisor),
        "pico": _card(pico),
        "fraca": _card(fraca),
        "expediente": [ln for ln in linhas if ln["expediente"]],
        "fora": [ln for ln in linhas if not ln["expediente"] and ln["total_bruto"] > 0],
        "linhas": linhas,
        "mapa": mapa,
        "detalhe": detalhe,
        "vazio": grand_n == 0,
    }


def carregar_vendas_intervalo(d0: date, d1: date) -> list[dict]:
    """Vendas do PDV no intervalo, sem devolução total. Entrega = pedido não cancelado."""
    from produtos.models import PedidoEntrega, VendaAgro

    desde, ate = _aware(d0, d1)
    rows = list(
        VendaAgro.objects.filter(
            devolvida_em__isnull=True,
            criado_em__gte=desde,
            criado_em__lte=ate,
        ).values(
            "id",
            "criado_em",
            "deposito",
            "total",
            "usuario_registro",
            "cliente_nome",
        )
    )
    ids = [r["id"] for r in rows]
    entregas: set[int] = set()
    for i in range(0, len(ids), 2000):
        bloco = ids[i : i + 2000]
        entregas.update(
            PedidoEntrega.objects.filter(venda_agro_id__in=bloco)
            .exclude(status=PedidoEntrega.Status.CANCELADO)
            .values_list("venda_agro_id", flat=True)
        )
    return [
        {
            "id": r["id"],
            "quando": r["criado_em"],
            "deposito": r["deposito"] or "",
            "total": r["total"] or 0,
            "operador": r["usuario_registro"] or "",
            "cliente": r["cliente_nome"] or "",
            "entrega": r["id"] in entregas,
        }
        for r in rows
    ]


def meta_card_hora(
    vendido,
    d0: date,
    d1: date,
    loja: str,
) -> dict[str, Any] | None:
    """Vendido deste relatório × meta do dia/período (mesma média da tela Meta)."""
    from produtos.meta_vendas_util import meta_fmt_moeda, meta_fmt_pct
    from produtos.vendas_lojas_util import (
        vendas_lojas_cmp_meta,
        vendas_lojas_cmp_meta_agora,
        vendas_lojas_meta_c_modos,
    )

    dep = loja if loja in ("centro", "vila") else None
    hoje = timezone.localdate()
    agora = timezone.localtime()
    dia_todo, ate_agora, mostra = vendas_lojas_meta_c_modos(
        d0, d1, dep, hoje=hoje, agora=agora
    )
    if dia_todo <= 0:
        return None
    vd = Decimal(str(vendido or 0)).quantize(Decimal("0.01"))
    if mostra:
        cmp = vendas_lojas_cmp_meta_agora(vd, ate_agora, dia_todo)
        ritmo = ate_agora
        titulo = "Até agora, no ritmo da meta"
    else:
        cmp = vendas_lojas_cmp_meta(vd, dia_todo)
        ritmo = dia_todo
        titulo = "Contra a meta do período"
    pct_dia = (vd / dia_todo * Decimal("100")).quantize(Decimal("0.1"))
    sentido = cmp.get("sentido") or "sem"
    return {
        "ok": True,
        "titulo": titulo,
        "mostra_ritmo": bool(mostra),
        "vendido_fmt": meta_fmt_moeda(vd),
        "ritmo_fmt": meta_fmt_moeda(ritmo),
        "meta_dia_fmt": meta_fmt_moeda(dia_todo),
        "pct_ritmo_fmt": meta_fmt_pct(cmp.get("pct")) if cmp.get("pct") is not None else "—",
        "pct_dia_fmt": meta_fmt_pct(pct_dia),
        "sentido": sentido,
        "acima": sentido == "acima",
        "abaixo": sentido == "abaixo",
    }


def montar_relatorio(request) -> tuple[dict, dict]:
    filtros = parse_periodo_hora(request)
    dias = listar_dias(filtros["d0"], filtros["d1"])
    base_dias = dias_costume(dias)
    loja = parse_loja_hora(request)
    canal = parse_canal_hora(request)
    visao = parse_visao_hora(request, len(dias))
    hora = parse_hora_detalhe(request)
    wd = parse_wd_detalhe(request)
    vendas = carregar_vendas_intervalo(filtros["d0"], filtros["d1"])
    baseline: list[dict] = []
    if base_dias:
        baseline = carregar_vendas_intervalo(min(base_dias), max(base_dias))
    grade = agregar_hora(
        vendas,
        baseline,
        dias=dias,
        dias_base=base_dias,
        loja=loja,
        canal=canal,
        visao=visao,
        hora_detalhe=hora,
        wd_detalhe=wd,
    )
    if canal == "todos":
        try:
            grade["meta"] = meta_card_hora(grade["total_bruto"], filtros["d0"], filtros["d1"], loja)
        except Exception:
            logger.exception("relatorio hora meta")
            grade["meta"] = None
            grade["meta_aviso"] = "Meta indisponível agora. O restante do relatório está valendo."
    else:
        grade["meta"] = None
        grade["meta_aviso"] = "A meta é da loja inteira. Para ver, deixe Balcão + entrega."
    filtros["loja"] = loja
    filtros["canal"] = canal
    filtros["visao"] = grade["visao"]
    filtros["n_dias"] = grade["n_dias"]
    filtros["hora"] = hora
    filtros["wd"] = wd
    return filtros, grade


def _estilo_cabeca(ws, headers: list[str]) -> None:
    fill = PatternFill("solid", fgColor="1E293B")
    font_h = Font(bold=True, color="FFFFFF")
    start = ws.max_row
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(row=start, column=col)
        cell.fill = fill
        cell.font = font_h
        cell.alignment = Alignment(horizontal="center")
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16


def montar_xlsx_hora(grade: dict, subtitulo: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hora a hora"
    ws.append(["Venda hora a hora"])
    ws["A1"].font = Font(bold=True, size=14)
    if subtitulo:
        ws.append([subtitulo])
    ws.append([grade.get("aviso_visao") or ""])
    ws.append([])
    headers = [
        "Hora",
        "Centro R$",
        "Vendas Centro",
        "Vila R$",
        "Vendas Vila",
        "Total R$",
        "Vendas",
        "Ticket",
        "% do dia",
        "Balcão R$",
        "Entrega R$",
        "Costume R$",
        "Variação %",
        "Quem atendeu",
    ]
    ws.append(headers)
    _estilo_cabeca(ws, headers)
    for ln in grade.get("linhas") or []:
        if not ln["expediente"] and ln["total_bruto"] <= 0:
            continue
        ws.append(
            [
                ln["rotulo"],
                ln["centro"],
                ln["n_centro"],
                ln["vila"],
                ln["n_vila"],
                ln["total"],
                ln["n"],
                ln["ticket"],
                ln["pct"],
                ln["balcao"],
                ln["entrega"],
                ln["base"] if ln["base"] is not None else "",
                ln["var"] if ln["var"] is not None else "",
                ln["quem"],
            ]
        )

    mp = wb.create_sheet("Mapa da semana")
    mp.append(["Mapa — média da hora naquele dia da semana"])
    mp["A1"].font = Font(bold=True, size=14)
    mp.append([])
    cab = ["Dia"] + [f"{h}h" for h in range(HORA_INI, HORA_FIM + 1)]
    mp.append(cab)
    _estilo_cabeca(mp, cab)
    for linha in grade.get("mapa") or []:
        vals = [linha["rotulo"]]
        for cel in linha["celulas"]:
            vals.append("" if cel.get("vazio") else cel.get("valor"))
        mp.append(vals)

    det = grade.get("detalhe")
    if det and det.get("vendas"):
        sh = wb.create_sheet("Vendas da hora")
        sh.append([f"Vendas — {det.get('rotulo')}"])
        sh["A1"].font = Font(bold=True, size=14)
        sh.append([])
        cab_d = ["Quando", "Loja", "Cliente", "Quem", "Tipo", "Total R$"]
        sh.append(cab_d)
        _estilo_cabeca(sh, cab_d)
        for v in det["vendas"]:
            sh.append(
                [v["quando_fmt"], v["loja"], v["cliente"], v["operador"], v["canal"], v["total"]]
            )

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def xlsx_hora_response(grade: dict, subtitulo: str):
    return xlsx_http_response("venda-hora-a-hora.xlsx", montar_xlsx_hora(grade, subtitulo))
